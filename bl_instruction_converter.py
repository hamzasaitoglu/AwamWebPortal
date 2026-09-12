import streamlit as st
import pandas as pd
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def extract_text_from_file(uploaded_file):
    """Extract text content from uploaded instruction file."""
    text = ""
    try:
        file_type = uploaded_file.name.split('.')[-1].lower()
        if file_type == 'pdf':
            import pypdf
            reader = pypdf.PdfReader(uploaded_file)
            for page in reader.pages:
                text += (page.extract_text() or "") + "\n"
        elif file_type in ['txt', 'csv']:
            text = uploaded_file.read().decode("utf-8", errors="ignore")
        else:
            text = "File uploaded successfully. Please review or fill fields manually."
    except Exception as e:
        text = f"Could not extract text automatically: {str(e)}"
    return text

def generate_bl_excel(booking_data, shipper_data, cnee_data, notify_data, containers_df):
    """Generate Excel file formatted according to Awam Logistics official template."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Shipping Instruction"
    ws.views.sheetView[0].showGridLines = True

    # Styling and Colors
    navy_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    white_bold_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    section_font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
    bold_font = Font(name="Calibri", size=10, bold=True)
    italic_font = Font(name="Calibri", size=10, italic=True)
    regular_font = Font(name="Calibri", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # 1. Title Banner (Row 2)
    ws.merge_cells("A2:G2")
    cell_a2 = ws["A2"]
    cell_a2.value = "Official Shipping Instruction Document | www.awamlogistics.com"
    cell_a2.fill = navy_fill
    cell_a2.font = white_bold_font
    cell_a2.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 24

    # 2. General Booking Info (Rows 4-9)
    b_fields = [
        ("Booking No:", booking_data.get("booking_no", "")),
        ("Shipping Line:", booking_data.get("line", "")),
        ("Vessel & Voyage:", booking_data.get("vessel_voyage", "")),
        ("POL (Loading Port):", booking_data.get("pol", "")),
        ("POD (Discharge Port):", booking_data.get("pod", "")),
        ("Freight Terms:", booking_data.get("freight_terms", "FREIGHT PREPAID"))
    ]

    for idx, (label, val) in enumerate(b_fields, start=4):
        ws.cell(row=idx, column=1, value=label).font = bold_font
        ws.cell(row=idx, column=3, value=val).font = regular_font

    # 3. Party Section Writer
    def write_party_section(start_row, section_title, party_data):
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=7)
        sec_cell = ws.cell(row=start_row, column=1, value=section_title)
        sec_cell.fill = section_fill
        sec_cell.font = section_font
        
        fields = [
            ("Company Name", party_data.get("name", "")),
            ("Address", party_data.get("address", "")),
            ("Tax Number", party_data.get("vat_no", "")),
            ("Tel", party_data.get("tel", "")),
            ("Email", party_data.get("email", ""))
        ]
        
        for offset, (lbl, val) in enumerate(fields, start=1):
            curr_row = start_row + offset
            ws.cell(row=curr_row, column=1, value=lbl).font = italic_font
            ws.cell(row=curr_row, column=3, value=val).font = regular_font

    write_party_section(11, "1. SHIPPER DETAILS", shipper_data)
    write_party_section(17, "2. CONSIGNEE DETAILS", cnee_data)
    write_party_section(23, "3. NOTIFY PARTY DETAILS", notify_data)

    # 4. Containers Table Header (Row 29)
    headers = ["Container No", "Seal No", "Type / HS Code", "Packages", "Description of Goods", "Gross Weight (KG)", "Volume (CBM)"]
    ws.row_dimensions[29].height = 22
    for col_num, h_text in enumerate(headers, 1):
        c = ws.cell(row=29, column=col_num, value=h_text)
        c.fill = navy_fill
        c.font = white_bold_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 5. Containers Data Rows (Row 30+)
    start_c_row = 30
    total_pkgs = 0
    total_gw = 0.0
    total_cbm = 0.0

    for i, row in containers_df.iterrows():
        r = start_c_row + i
        ws.cell(row=r, column=1, value=row.get("Container No", "")).font = regular_font
        ws.cell(row=r, column=2, value=row.get("Seal No", "")).font = regular_font
        ws.cell(row=r, column=3, value=row.get("HS Code / Type", "")).font = regular_font
        
        # Packages
        pkg_val = row.get("Packages", "")
        ws.cell(row=r, column=4, value=pkg_val).font = regular_font
        try:
            total_pkgs += int(pkg_val)
        except (ValueError, TypeError):
            pass

        ws.cell(row=r, column=5, value=row.get("Description", "")).font = regular_font
        
        # Gross Weight
        gw_val = float(row.get("Gross Weight (KG)", 0) or 0)
        gw_cell = ws.cell(row=r, column=6, value=gw_val)
        gw_cell.number_format = '#,##0.00'
        gw_cell.font = regular_font
        total_gw += gw_val
        
        # Volume (CBM)
        cbm_val = float(row.get("Volume (CBM)", 0) or 0)
        cbm_cell = ws.cell(row=r, column=7, value=cbm_val)
        cbm_cell.number_format = '#,##0.00'
        cbm_cell.font = regular_font
        total_cbm += cbm_val

        for col_num in range(1, 8):
            ws.cell(row=r, column=col_num).border = thin_border

    # 6. TOTAL Row Addition
    tot_row = start_c_row + len(containers_df)
    ws.cell(row=tot_row, column=1, value="TOTAL").font = bold_font
    ws.cell(row=tot_row, column=4, value=total_pkgs if total_pkgs > 0 else "").font = bold_font
    
    tot_gw_cell = ws.cell(row=tot_row, column=6, value=total_gw)
    tot_gw_cell.number_format = '#,##0.00'
    tot_gw_cell.font = bold_font

    tot_cbm_cell = ws.cell(row=tot_row, column=7, value=total_cbm)
    tot_cbm_cell.number_format = '#,##0.00'
    tot_cbm_cell.font = bold_font

    for col_num in range(1, 8):
        c = ws.cell(row=tot_row, column=col_num)
        c.fill = section_fill
        c.border = thin_border

    # Adjust Column Widths (Enlarged Description Column to 60)
    col_widths = {1: 20, 2: 16, 3: 18, 4: 15, 5: 60, 6: 20, 7: 16}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def render_bl_instruction_converter():
    st.title("📦 B/L Instruction Converter")
    st.markdown("---")

    # 1. Upload Section
    st.subheader("1. Upload Shipper Instructions")
    uploaded_file = st.file_uploader("Upload Instructions File (PDF, TXT, CSV, XLSX)", type=["pdf", "txt", "csv", "xlsx"])
    
    if uploaded_file is not None:
        extracted_text = extract_text_from_file(uploaded_file)
        with st.expander("Preview Extracted Text", expanded=False):
            st.text_area("Extracted Content:", extracted_text, height=150)

    # 2. Booking Details
    st.subheader("2. Booking & Transport Details")
    col_b1, col_b2, col_b3 = st.columns(3)
    with col_b1:
        booking_no = st.text_input("BOOKING NO")
        pol = st.text_input("POL (Port of Loading)")
    with col_b2:
        line = st.text_input("Shipping Line")
        pod = st.text_input("POD (Port of Discharge)")
    with col_b3:
        vessel_voyage = st.text_input("Vessel & Voyage")
        freight_terms = st.selectbox("Freight Terms", ["FREIGHT PREPAID", "FREIGHT COLLECT"])

    booking_data = {
        "booking_no": booking_no, "line": line, "vessel_voyage": vessel_voyage,
        "pol": pol, "pod": pod, "freight_terms": freight_terms
    }

    # 3. Parties Details (Single Page Layout with Dynamic Copy from Consignee to Notify)
    st.subheader("3. Parties Details")
    col_p1, col_p2, col_p3 = st.columns(3, gap="medium")

    with col_p1:
        st.markdown("##### 🏢 Shipper Details")
        s_name = st.text_input("Shipper Name", key="s_name")
        s_addr = st.text_area("Shipper Address", key="s_addr", height=100)
        s_tel = st.text_input("Shipper Tel", key="s_tel")
        s_email = st.text_input("Shipper Email", key="s_email")
        s_vat = st.text_input("Shipper VAT / Tax No", key="s_vat")

    with col_p2:
        st.markdown("##### 🏬 Consignee Details (CNEE)")
        c_name = st.text_input("Consignee Name", key="c_name")
        c_addr = st.text_area("Consignee Address", key="c_addr", height=100)
        c_tel = st.text_input("Consignee Tel", key="c_tel")
        c_email = st.text_input("Consignee Email", key="c_email")
        c_vat = st.text_input("Consignee VAT / Tax No", key="c_vat")

    with col_p3:
        st.markdown("##### 🔔 Notify Party Details")
        same_as_cnee = st.checkbox("Same as Consignee (CNEE)", key="same_as_cnee")
        
        if same_as_cnee:
            n_name = st.text_input("Notify Name", value=c_name, key="n_name_dis", disabled=True)
            n_addr = st.text_area("Notify Address", value=c_addr, key="n_addr_dis", height=100, disabled=True)
            n_tel = st.text_input("Notify Tel", value=c_tel, key="n_tel_dis", disabled=True)
            n_email = st.text_input("Notify Email", value=c_email, key="n_email_dis", disabled=True)
            n_vat = st.text_input("Notify VAT / Tax No", value=c_vat, key="n_vat_dis", disabled=True)
        else:
            n_name = st.text_input("Notify Name", key="n_name")
            n_addr = st.text_area("Notify Address", key="n_addr", height=100)
            n_tel = st.text_input("Notify Tel", key="n_tel")
            n_email = st.text_input("Notify Email", key="n_email")
            n_vat = st.text_input("Notify VAT / Tax No", key="n_vat")

    shipper_data = {"name": s_name, "address": s_addr, "tel": s_tel, "email": s_email, "vat_no": s_vat}
    cnee_data = {"name": c_name, "address": c_addr, "tel": c_tel, "email": c_email, "vat_no": c_vat}
    notify_data = {"name": n_name, "address": n_addr, "tel": n_tel, "email": n_email, "vat_no": n_vat}

    # 4. Containers Details Table with Summary Bar
    st.subheader("4. Containers & Cargo Details")
    
    if 'containers_data' not in st.session_state:
        st.session_state.containers_data = pd.DataFrame([
            {"Container No": "", "Seal No": "", "HS Code / Type": "40' HC", "Packages": "", "Description": "", "Gross Weight (KG)": 0.0, "Volume (CBM)": 0.0}
        ])

    edited_df = st.data_editor(
        st.session_state.containers_data,
        num_rows="dynamic",
        use_container_width=True,
        column_config={
            "Container No": st.column_config.TextColumn("Container No", required=True, width="medium"),
            "Seal No": st.column_config.TextColumn("Seal No", width="small"),
            "HS Code / Type": st.column_config.TextColumn("HS Code / Type", width="medium"),
            "Packages": st.column_config.TextColumn("Packages", width="small"),
            "Description": st.column_config.TextColumn("Description", width="large"),
            "Gross Weight (KG)": st.column_config.NumberColumn("Gross Weight (KG)", format="%.2f", width="medium"),
            "Volume (CBM)": st.column_config.NumberColumn("Volume (CBM)", format="%.2f", width="medium")
        },
        key="container_editor"
    )

    # Calculate UI Totals
    total_containers = len(edited_df)
    total_gw = pd.to_numeric(edited_df["Gross Weight (KG)"], errors='coerce').sum()
    total_cbm = pd.to_numeric(edited_df["Volume (CBM)"], errors='coerce').sum()

    # Display Summary Bar Below Table
    st.markdown("**Summary Totals:**")
    m1, m2, m3 = st.columns(3)
    m1.metric("Total Containers", f"{total_containers}")
    m2.metric("Total Gross Weight", f"{total_gw:,.2f} KG")
    m3.metric("Total Volume", f"{total_cbm:,.2f} CBM")

    st.markdown("---")
    
    # 5. Generate and Download
    if st.button("Generate Formatted Excel File", type="primary"):
        excel_file = generate_bl_excel(booking_data, shipper_data, cnee_data, notify_data, edited_df)
        file_filename = f"Shipping_Instruction_{booking_no if booking_no else 'Awam'}.xlsx"
        
        st.download_button(
            label="📥 Download Excel File",
            data=excel_file,
            file_name=file_filename,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

if __name__ == "__main__":
    render_bl_instruction_converter()